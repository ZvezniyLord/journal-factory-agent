from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def _norm(value: Any) -> str:
    text = "" if value is None else str(value)
    text = text.strip().lower().replace("’", "'")
    return re.sub(r"\s+", " ", text)


COLUMN_ALIASES = {
    "id": ("id", "№", "номер", "код", "ідентифікатор"),
    "authors": ("автор", "автори", "піб", "имя", "ім'я", "прізвище ім'я", "participant", "author"),
    "title": ("назва", "назва статті", "назва тез", "тема", "title", "article title"),
    "section": ("секція", "розділ", "section"),
    "free_listener": ("вільний слухач", "слухач", "free listener"),
    "doi": ("doi", "дой"),
    "printed_copy": ("друк", "друкований примірник", "printed", "printed copy"),
}


def guess_columns(headers: list[Any]) -> dict[str, int]:
    normalized = [_norm(v) for v in headers]
    result: dict[str, int] = {}
    for semantic, aliases in COLUMN_ALIASES.items():
        for index, header in enumerate(normalized):
            if not header:
                continue
            if any(alias == header or alias in header for alias in aliases):
                result.setdefault(semantic, index)
                break
    return result


def inspect_excel_registry(
    path: str | Path,
    *,
    sheet_name: str | None = None,
    header_row: int = 1,
) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]

    rows = list(sheet.iter_rows(values_only=True))
    if len(rows) < header_row:
        return {
            "path": str(Path(path)),
            "sheet": sheet.title,
            "headers": [],
            "column_map": {},
            "rows": [],
        }

    headers = list(rows[header_row - 1])
    column_map = guess_columns(headers)
    output_rows: list[dict[str, Any]] = []

    for excel_row_number, values in enumerate(rows[header_row:], start=header_row + 1):
        if not any(value not in (None, "") for value in values):
            continue

        raw = {
            str(headers[i] if i < len(headers) and headers[i] not in (None, "") else f"column_{i+1}"): value
            for i, value in enumerate(values)
            if value not in (None, "")
        }
        semantic: dict[str, Any] = {}
        for key, index in column_map.items():
            if index < len(values):
                semantic[key] = values[index]

        output_rows.append(
            {
                "excel_row": excel_row_number,
                "semantic": semantic,
                "raw": raw,
            }
        )

    return {
        "path": str(Path(path)),
        "sheet": sheet.title,
        "headers": headers,
        "column_map": column_map,
        "row_count": len(output_rows),
        "rows": output_rows,
    }
