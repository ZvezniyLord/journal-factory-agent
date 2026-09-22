from __future__ import annotations

import base64
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from .reference_numbering import (
    bind_reference_numbering,
    create_article_reference_num,
    ensure_shared_reference_abstract_num,
)


_TINY_PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAusB9Wl2s7sAAAAASUVORK5CYII="
)


def _ensure_style(doc: Document, name: str, *, size_pt: float = 11.0):
    styles = doc.styles
    try:
        style = styles[name]
    except KeyError:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    style.font.name = "Times New Roman"
    style.font.size = Pt(size_pt)
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Times New Roman")
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Times New Roman")
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Times New Roman")
    style._element.get_or_add_rPr().rFonts.set(qn("w:cs"), "Times New Roman")
    for node in style._element.findall(qn("w:autoRedefine")):
        style._element.remove(node)
    return style


def _configure_body_style(style) -> None:
    pf = style.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.0
    pf.space_after = Pt(0)
    pf.first_line_indent = Cm(1)


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    rpr.extend([color, underline])
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.extend([rpr, text_node])
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def create_golden_fixture(path: str | Path) -> Path:
    output = Path(path)
    output.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    body = _ensure_style(doc, "ASTRA_BODY")
    _configure_body_style(body)
    refer = _ensure_style(doc, "REFER")
    ref_title = _ensure_style(doc, "REF-TITLE")
    section = _ensure_style(doc, "SECTION", size_pt=24)
    reference_plan = ensure_shared_reference_abstract_num(doc)

    p = doc.add_paragraph("FRONT MATTER SENTINEL — MUST REMAIN UNCHANGED")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()
    doc.add_paragraph("TABLE OF CONTENTS")
    doc.add_page_break()

    for article_index in (1, 2):
        doc.add_paragraph(f"SECTION {article_index}", style=section)
        doc.add_paragraph(f"UDC 001.{article_index}")
        doc.add_paragraph(f"Author {article_index}")
        doc.add_paragraph(f"ARTICLE TITLE {article_index}")

        p = doc.add_paragraph(style=body)
        p.add_run("Normal ")
        r = p.add_run("bold")
        r.bold = True
        p.add_run(" ")
        r = p.add_run("italic")
        r.italic = True
        p.add_run(" ")
        r = p.add_run("underline")
        r.underline = True
        p.add_run(" H")
        r = p.add_run("2")
        r.font.subscript = True
        p.add_run("O x")
        r = p.add_run("2")
        r.font.superscript = True

        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "A"
        table.cell(0, 1).text = "B"
        table.cell(1, 0).text = "1"
        table.cell(1, 1).text = "2"

        link_p = doc.add_paragraph(style=body)
        _add_hyperlink(link_p, "Example link", "https://example.com")

        with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
            tmp.write(_TINY_PNG)
            png_path = Path(tmp.name)
        try:
            doc.add_picture(str(png_path), width=Cm(0.5))
        finally:
            png_path.unlink(missing_ok=True)

        doc.add_paragraph("REFERENCES", style=ref_title)
        num_id = create_article_reference_num(doc, reference_plan)
        for ref_index in range(1, 3):
            rp = doc.add_paragraph(
                f"Reference {article_index}.{ref_index}",
                style=refer,
            )
            bind_reference_numbering(rp, num_id)

        continuation = doc.add_paragraph(
            f"https://doi.org/10.0000/example-{article_index}",
            style=refer,
        )
        # No numPr: this is an explicit continuation line.

        if article_index == 1:
            doc.add_page_break()

    doc.save(output)
    return output
