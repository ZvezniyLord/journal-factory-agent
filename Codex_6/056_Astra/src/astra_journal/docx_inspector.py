from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any
from zipfile import ZipFile

from docx import Document


def _run_record(run) -> dict[str, Any]:
    font = run.font
    return {
        "text": run.text,
        "bold": run.bold,
        "italic": run.italic,
        "underline": bool(run.underline) if run.underline is not None else None,
        "superscript": font.superscript,
        "subscript": font.subscript,
    }


def _paragraph_record(paragraph, index: int) -> dict[str, Any]:
    ppr = paragraph._p.pPr
    num_id = None
    ilvl = None
    if ppr is not None and ppr.numPr is not None:
        num = ppr.numPr.numId
        level = ppr.numPr.ilvl
        num_id = num.val if num is not None else None
        ilvl = level.val if level is not None else None

    return {
        "index": index,
        "text": paragraph.text,
        "style_id": paragraph.style.style_id if paragraph.style is not None else None,
        "style_name": paragraph.style.name if paragraph.style is not None else None,
        "num_id": int(num_id) if num_id is not None else None,
        "ilvl": int(ilvl) if ilvl is not None else None,
        "runs": [_run_record(run) for run in paragraph.runs if run.text],
    }


def inspect_docx_source(path: str | Path, *, excerpt_paragraphs: int = 80) -> dict[str, Any]:
    file_path = Path(path)
    document = Document(file_path)

    paragraphs = [
        _paragraph_record(paragraph, index)
        for index, paragraph in enumerate(document.paragraphs)
        if paragraph.text.strip()
    ]
    full_text = "\n".join(p["text"] for p in paragraphs)

    with ZipFile(file_path) as archive:
        names = archive.namelist()
        media = [name for name in names if name.startswith("word/media/")]
        document_xml = archive.read("word/document.xml")
        hyperlink_count = document_xml.count(b"<w:hyperlink")

    return {
        "path": str(file_path),
        "sha256": hashlib.sha256(file_path.read_bytes()).hexdigest(),
        "paragraph_count": len(paragraphs),
        "table_count": len(document.tables),
        "media_count": len(media),
        "hyperlink_count": hyperlink_count,
        "text_sha256": hashlib.sha256(full_text.encode("utf-8")).hexdigest(),
        "excerpt": paragraphs[:excerpt_paragraphs],
    }
